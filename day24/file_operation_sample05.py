"""
Version: 1.0
Author: ChatGPT
Descrtiption: 生成统计图表的 Excel 写入示例代码。
Date: 2025-10-30

Keyword arguments:
argument -- description
Return: return_description
"""
import openpyxl
from openpyxl.chart import BarChart, Reference

data = [
    ['月份', '销售额'],
    ['一月', 15000],
    ['二月', 18000],
    ['三月', 22000],
    ['四月', 24000],
    ['五月', 20000],
    ['六月', 26000],
    ['七月', 30000],
    ['八月', 28000],
    ['九月', 32000],
    ['十月', 35000],
    ['十一月', 40000],
    ['十二月', 45000],
]

# 创建工作簿和工作表
workbook = openpyxl.Workbook()
sheet = workbook.active
sheet.title = '销售数据'
# 写入数据
for row_index, row_data in enumerate(data, start=1):
    for col_index, cell_value in enumerate(row_data, start=1):
        sheet.cell(row=row_index, column=col_index, value=cell_value)   
# 创建柱状图
chart = BarChart()
chart.title = '月度销售额统计'
chart.x_axis.title = '月份'
chart.y_axis.title = '销售额 (元)'
data_ref = Reference(sheet, min_col=2, min_row=1, max_row=len(data))
cats_ref = Reference(sheet, min_col=1, min_row=2, max_row=len(data))
chart.add_data(data_ref, titles_from_data=True)
chart.set_categories(cats_ref)
# 将图表添加到工作表
sheet.add_chart(chart, 'E5')
# 保存工作簿
workbook.save('销售数据统计表.xlsx')

